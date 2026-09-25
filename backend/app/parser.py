import os
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Tuple
import pandas as pd

from docling.datamodel.base_models import InputFormat
from docling.document_converter import DocumentConverter, PdfFormatOption
from docling.datamodel.pipeline_options import PdfPipelineOptions
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

from app.config import settings

class DocumentParser:
    """
    Enterprise multi-format document parser.
    Specialized for dense financial and legal layouts (PDFs, DOCX, XLSX, CSV, TXT).
    """

    def __init__(self):
        # Configure Docling with optimized CPU pipeline options
        # Table structure extraction enabled, heavy OCR disabled for fast offline processing
        pipeline_options = PdfPipelineOptions()
        pipeline_options.do_ocr = False
        pipeline_options.do_table_structure = True

        self.docling_converter = DocumentConverter(
            format_options={
                InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)
            }
        )
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            separators=["\n## ", "\n### ", "\n\n", "\n|", "\n", " ", ""]
        )

    def parse_file(self, file_path: str, original_filename: str) -> Tuple[str, Dict[str, Any]]:
        """
        Parses multi-format document into structured markdown text and metadata.
        """
        path = Path(file_path)
        ext = path.suffix.lower() if path.suffix else Path(original_filename).suffix.lower()
        doc_id = str(uuid.uuid4())[:8]

        metadata: Dict[str, Any] = {
            "doc_id": doc_id,
            "filename": original_filename,
            "file_type": ext,
            "created_at": datetime.now().isoformat(),
            "byte_size": os.path.getsize(file_path) if os.path.exists(file_path) else 0
        }

        extracted_markdown = ""

        # 1. Financial & Legal PDFs via Docling
        if ext == ".pdf":
            try:
                result = self.docling_converter.convert(file_path)
                extracted_markdown = result.document.export_to_markdown()
                metadata["parser"] = "docling_pdf"
            except Exception:
                # Fallback to pypdfium2 if docling encounters specialized PDF errors
                metadata["parser"] = "pypdfium2_fallback"
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(file_path)
                text_pages = [page.get_textpage().get_text_range() for page in pdf]
                extracted_markdown = "\n\n".join(text_pages)

        # 2. Financial Spreadsheets (XLSX, XLS, CSV)
        elif ext in [".xlsx", ".xls"]:
            metadata["parser"] = "pandas_excel"
            excel_file = pd.ExcelFile(file_path)
            sheets_md = []
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheets_md.append(f"### Sheet: {sheet_name}\n" + df.to_markdown(index=False))
            extracted_markdown = "\n\n".join(sheets_md)

        elif ext == ".csv":
            metadata["parser"] = "pandas_csv"
            df = pd.read_csv(file_path)
            extracted_markdown = df.to_markdown(index=False)

        # 3. Word Documents (DOCX)
        elif ext == ".docx":
            try:
                result = self.docling_converter.convert(file_path)
                extracted_markdown = result.document.export_to_markdown()
                metadata["parser"] = "docling_docx"
            except Exception:
                import docx
                doc = docx.Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                extracted_markdown = "\n\n".join(paragraphs)
                metadata["parser"] = "python_docx_fallback"

        # 4. Plain Text & Markdown
        elif ext in [".txt", ".md"]:
            metadata["parser"] = "text_reader"
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    extracted_markdown = f.read()
            except UnicodeDecodeError:
                with open(file_path, "r", encoding="latin-1") as f:
                    extracted_markdown = f.read()
        else:
            raise ValueError(f"Unsupported file format: {ext}")

        metadata["char_count"] = len(extracted_markdown)
        metadata["has_tables"] = "|" in extracted_markdown
        return extracted_markdown, metadata

    def chunk_document(self, text: str, metadata: Dict[str, Any]) -> List[Document]:
        """
        Splits extracted markdown text into semantically cohesive chunks with rich metadata.
        """
        raw_chunks = self.text_splitter.split_text(text)
        total_chunks = len(raw_chunks)
        documents = []

        for idx, chunk_text in enumerate(raw_chunks):
            chunk_metadata = {
                "doc_id": metadata["doc_id"],
                "source": metadata["filename"],
                "file_type": metadata["file_type"],
                "chunk_index": idx,
                "total_chunks": total_chunks,
                "char_length": len(chunk_text)
            }
            documents.append(
                Document(page_content=chunk_text, metadata=chunk_metadata)
            )

        return documents

# Singleton parser instance
parser_engine = DocumentParser()
